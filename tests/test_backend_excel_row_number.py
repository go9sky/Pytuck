"""
Excel 后端行号映射功能测试

测试 Excel 后端的 row_number_mapping 选项：
- 外部 Excel 文件读取（无 _pytuck_tables schema）
- 行号作为主键 (as_pk)
- 行号映射到字段 (field)
- 行号持久化 (persist_row_number)
- 覆盖选项 (row_number_override)
"""

from pathlib import Path
from typing import Type

import pytest

from pytuck import Storage, Session, Column, PureBaseModel, declarative_base
from pytuck import select, insert
from pytuck.common.options import ExcelBackendOptions


def create_external_excel(file_path: Path, sheet_name: str = 'Sheet1') -> None:
    """创建一个不带 Pytuck 元数据的外部 Excel 文件"""
    try:
        from openpyxl import Workbook
    except ImportError:
        pytest.skip("openpyxl not installed")

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # 写入表头和数据
    ws.append(['name', 'age', 'city'])
    ws.append(['Alice', 25, 'Beijing'])
    ws.append(['Bob', 30, 'Shanghai'])
    ws.append(['Charlie', 35, 'Guangzhou'])

    wb.save(str(file_path))


def create_external_excel_with_row_num_field(file_path: Path) -> None:
    """创建一个带有 row_num 字段的外部 Excel 文件"""
    try:
        from openpyxl import Workbook
    except ImportError:
        pytest.skip("openpyxl not installed")

    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'

    # 写入表头和数据（包含 row_num 字段）
    ws.append(['name', 'row_num', 'value'])
    ws.append(['Alice', 100, 'a'])  # row_num 有值
    ws.append(['Bob', None, 'b'])   # row_num 为空
    ws.append(['Charlie', 200, 'c'])  # row_num 有值

    wb.save(str(file_path))


class TestExternalExcelDefault:
    """测试外部 Excel 默认行为（无映射）"""

    def test_load_external_excel_default(self, tmp_path: Path) -> None:
        """外部 Excel，mapping=None（默认），应使用自增主键"""
        excel_file = tmp_path / 'external.xlsx'
        create_external_excel(excel_file)

        db = Storage(file_path=str(excel_file), engine='excel')

        # 验证表加载成功
        assert 'Sheet1' in db.tables
        table = db.tables['Sheet1']

        # 验证数据
        assert len(table.data) == 3

        # 验证主键是自增的 (1, 2, 3)
        assert set(table.data.keys()) == {1, 2, 3}

        db.close()


class TestRowNumberAsPrimaryKey:
    """测试行号作为主键"""

    def test_row_number_as_pk(self, tmp_path: Path) -> None:
        """外部 Excel，mapping='as_pk'，行号应作为主键"""
        excel_file = tmp_path / 'external.xlsx'
        create_external_excel(excel_file)

        opts = ExcelBackendOptions(row_number_mapping='as_pk')
        db = Storage(file_path=str(excel_file), engine='excel', backend_options=opts)

        table = db.tables['Sheet1']

        # 验证主键是 Excel 行号（2, 3, 4）
        # 第一行是表头，数据从第 2 行开始
        assert set(table.data.keys()) == {2, 3, 4}

        # 验证数据内容
        assert table.data[2]['name'] == 'Alice'
        assert table.data[3]['name'] == 'Bob'
        assert table.data[4]['name'] == 'Charlie'

        # 验证 next_id 更新
        assert table.next_id == 5

        db.close()

    def test_row_number_as_pk_with_session(self, tmp_path: Path) -> None:
        """使用 Session 查询行号主键的数据"""
        excel_file = tmp_path / 'external.xlsx'
        create_external_excel(excel_file)

        opts = ExcelBackendOptions(row_number_mapping='as_pk')
        db = Storage(file_path=str(excel_file), engine='excel', backend_options=opts)
        Base: Type[PureBaseModel] = declarative_base(db)

        class Person(Base):
            __tablename__ = 'Sheet1'
            id = Column(int, primary_key=True)
            name = Column(str)
            age = Column(str)  # 外部 Excel 默认都是 str
            city = Column(str)

        session = Session(db)

        # 按主键查询（行号）
        result = session.execute(select(Person).where(Person.id == 2)).first()
        assert result is not None
        assert result.name == 'Alice'

        session.close()
        db.close()


class TestRowNumberAsField:
    """测试行号映射到字段"""

    def test_row_number_as_field(self, tmp_path: Path) -> None:
        """外部 Excel，mapping='field'，行号应映射到指定字段"""
        excel_file = tmp_path / 'external.xlsx'
        create_external_excel(excel_file)

        opts = ExcelBackendOptions(
            row_number_mapping='field',
            row_number_field_name='row_num'
        )
        db = Storage(file_path=str(excel_file), engine='excel', backend_options=opts)

        table = db.tables['Sheet1']

        # 验证主键是自增的（1, 2, 3）
        assert set(table.data.keys()) == {1, 2, 3}

        # 验证 row_num 字段包含 Excel 行号
        assert table.data[1]['row_num'] == 2
        assert table.data[2]['row_num'] == 3
        assert table.data[3]['row_num'] == 4

        # 验证 row_num 列存在于 columns 中
        assert 'row_num' in table.columns

        db.close()

    def test_row_number_field_custom_name(self, tmp_path: Path) -> None:
        """使用自定义字段名"""
        excel_file = tmp_path / 'external.xlsx'
        create_external_excel(excel_file)

        opts = ExcelBackendOptions(
            row_number_mapping='field',
            row_number_field_name='excel_line'
        )
        db = Storage(file_path=str(excel_file), engine='excel', backend_options=opts)

        table = db.tables['Sheet1']

        # 验证使用了自定义字段名
        assert 'excel_line' in table.columns
        assert table.data[1]['excel_line'] == 2

        db.close()


class TestFieldConflict:
    """测试字段名冲突"""

    def test_field_conflict_no_override(self, tmp_path: Path) -> None:
        """字段已存在且有值，override=False，不应覆盖"""
        excel_file = tmp_path / 'external.xlsx'
        create_external_excel_with_row_num_field(excel_file)

        opts = ExcelBackendOptions(
            row_number_mapping='field',
            row_number_field_name='row_num',
            row_number_override=False
        )
        db = Storage(file_path=str(excel_file), engine='excel', backend_options=opts)

        table = db.tables['Sheet1']

        # Alice 的 row_num 是 100（原有值，不覆盖）
        alice = next(r for r in table.data.values() if r['name'] == 'Alice')
        assert alice['row_num'] == 100

        # Bob 的 row_num 原来是 None，应填充行号
        bob = next(r for r in table.data.values() if r['name'] == 'Bob')
        assert bob['row_num'] == 3  # Excel 第 3 行

        # Charlie 的 row_num 是 200（原有值，不覆盖）
        charlie = next(r for r in table.data.values() if r['name'] == 'Charlie')
        assert charlie['row_num'] == 200

        db.close()

    def test_field_conflict_with_override(self, tmp_path: Path) -> None:
        """字段已存在且有值，override=True，应覆盖"""
        excel_file = tmp_path / 'external.xlsx'
        create_external_excel_with_row_num_field(excel_file)

        opts = ExcelBackendOptions(
            row_number_mapping='field',
            row_number_field_name='row_num',
            row_number_override=True
        )
        db = Storage(file_path=str(excel_file), engine='excel', backend_options=opts)

        table = db.tables['Sheet1']

        # 所有记录的 row_num 都应该是 Excel 行号
        alice = next(r for r in table.data.values() if r['name'] == 'Alice')
        assert alice['row_num'] == 2

        bob = next(r for r in table.data.values() if r['name'] == 'Bob')
        assert bob['row_num'] == 3

        charlie = next(r for r in table.data.values() if r['name'] == 'Charlie')
        assert charlie['row_num'] == 4

        db.close()


class TestPytuckExcelWithMapping:
    """测试 Pytuck 创建的 Excel 文件与行号映射"""

    def test_pytuck_excel_mapping_ignored_by_default(self, tmp_path: Path) -> None:
        """Pytuck 创建的文件，mapping='as_pk'，override=False，应忽略映射"""
        excel_file = tmp_path / 'pytuck.xlsx'

        # 先创建一个 Pytuck Excel 文件
        db1 = Storage(file_path=str(excel_file), engine='excel')
        Base1: Type[PureBaseModel] = declarative_base(db1)

        class User(Base1):
            __tablename__ = 'users'
            id = Column(int, primary_key=True)
            name = Column(str)

        session1 = Session(db1)
        session1.execute(insert(User).values(name='Alice'))
        session1.execute(insert(User).values(name='Bob'))
        session1.commit()
        db1.flush()
        db1.close()

        # 重新打开，带行号映射选项但 override=False
        opts = ExcelBackendOptions(
            row_number_mapping='as_pk',
            row_number_override=False
        )
        db2 = Storage(file_path=str(excel_file), engine='excel', backend_options=opts)

        table = db2.tables['users']

        # 应使用原有主键（1, 2），而非行号
        assert set(table.data.keys()) == {1, 2}

        db2.close()

    def test_pytuck_excel_mapping_with_override(self, tmp_path: Path) -> None:
        """Pytuck 创建的文件，mapping='as_pk'，override=True，应强制使用行号"""
        excel_file = tmp_path / 'pytuck.xlsx'

        # 先创建一个 Pytuck Excel 文件
        db1 = Storage(file_path=str(excel_file), engine='excel')
        Base1: Type[PureBaseModel] = declarative_base(db1)

        class User(Base1):
            __tablename__ = 'users'
            id = Column(int, primary_key=True)
            name = Column(str)

        session1 = Session(db1)
        session1.execute(insert(User).values(name='Alice'))
        session1.execute(insert(User).values(name='Bob'))
        session1.commit()
        db1.flush()
        db1.close()

        # 重新打开，带行号映射选项且 override=True
        opts = ExcelBackendOptions(
            row_number_mapping='as_pk',
            row_number_override=True
        )
        db2 = Storage(file_path=str(excel_file), engine='excel', backend_options=opts)

        table = db2.tables['users']

        # 应使用行号作为主键（2, 3）
        assert set(table.data.keys()) == {2, 3}

        db2.close()


class TestPersistRowNumber:
    """测试行号持久化"""

    def test_persist_row_number_on_save(self, tmp_path: Path) -> None:
        """保存时，persist_row_number=True，行号应写入文件"""
        excel_file = tmp_path / 'external.xlsx'
        create_external_excel(excel_file)

        # 加载并启用行号持久化
        opts = ExcelBackendOptions(
            row_number_mapping='field',
            row_number_field_name='row_num',
            persist_row_number=True
        )
        db = Storage(file_path=str(excel_file), engine='excel', backend_options=opts)

        # 验证加载后 row_num 列存在
        table = db.tables['Sheet1']
        assert 'row_num' in table.columns

        # 进行一些修改以触发 dirty 标记（添加新记录）
        Base: Type[PureBaseModel] = declarative_base(db)

        class Person(Base):
            __tablename__ = 'Sheet1'
            id = Column(int, primary_key=True)
            name = Column(str)
            age = Column(str)
            city = Column(str)
            row_num = Column(int)

        session = Session(db)
        session.execute(insert(Person).values(name='Dave', age='40', city='Shenzhen'))
        session.commit()

        # 保存（会覆盖原文件）
        db.flush()
        db.close()

        # 重新打开验证
        db2 = Storage(file_path=str(excel_file), engine='excel')

        table = db2.tables['Sheet1']

        # 验证 row_num 列存在并有正确的值
        assert 'row_num' in table.columns
        # 保存后重新加载，原有记录的 row_num 应该是 2, 3, 4
        # 新记录的 row_num 应该是 5（保存时的行号）
        row_nums = [r['row_num'] for r in table.data.values()]
        assert 2 in row_nums
        assert 3 in row_nums
        assert 4 in row_nums
        assert 5 in row_nums  # 新增记录

        db2.close()


class TestNextIdUpdate:
    """测试 next_id 更新"""

    def test_next_id_after_row_number_pk(self, tmp_path: Path) -> None:
        """行号作为主键后，next_id 应正确更新"""
        excel_file = tmp_path / 'external.xlsx'
        create_external_excel(excel_file)

        opts = ExcelBackendOptions(row_number_mapping='as_pk')
        db = Storage(file_path=str(excel_file), engine='excel', backend_options=opts)

        table = db.tables['Sheet1']

        # 数据有 3 行（行号 2, 3, 4），next_id 应为 5
        assert table.next_id == 5

        db.close()
